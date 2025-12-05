#!/usr/bin/env python3
"""
Test script to verify seasonal data appears in exports.
This is a verification script for task 6 of the seasonal-attributes spec.
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook

def test_json_export():
    """Verify JSON export includes seasonal boolean fields."""
    print("\n=== Testing JSON Export ===")
    json_path = Path("reports/games.json")
    
    if not json_path.exists():
        print(f"❌ JSON file not found at {json_path}")
        return False
    
    with open(json_path, 'r', encoding='utf-8') as f:
        games = json.load(f)
    
    if not games:
        print("❌ JSON file is empty")
        return False
    
    # Check first game has seasonal fields
    first_game = games[0]
    seasonal_fields = [
        'seasonal_spring',
        'seasonal_summer', 
        'seasonal_fall_halloween',
        'seasonal_winter_christmas'
    ]
    
    missing_fields = [field for field in seasonal_fields if field not in first_game]
    
    if missing_fields:
        print(f"❌ Missing seasonal fields in JSON: {missing_fields}")
        return False
    
    print(f"✓ JSON export contains all 4 seasonal fields")
    
    # Find games with seasonal tags
    seasonal_games = [g for g in games if any(g.get(f) for f in seasonal_fields)]
    print(f"✓ Found {len(seasonal_games)} games with seasonal tags out of {len(games)} total")
    
    # Show some examples
    if seasonal_games:
        print("\nExample seasonal games:")
        for game in seasonal_games[:3]:
            title = game.get('title', 'Unknown')
            seasons = [f.replace('seasonal_', '') for f in seasonal_fields if game.get(f)]
            print(f"  - {title}: {', '.join(seasons)}")
    
    return True

def test_excel_export():
    """Verify Excel export includes four seasonal columns."""
    print("\n=== Testing Excel Export ===")
    excel_path = Path("reports/games.xlsx")
    
    if not excel_path.exists():
        print(f"❌ Excel file not found at {excel_path}")
        return False
    
    wb = load_workbook(excel_path)
    ws = wb.active
    
    # Get headers from first row
    headers = [cell.value for cell in ws[1]]
    
    seasonal_columns = [
        'seasonal_spring',
        'seasonal_summer',
        'seasonal_fall_halloween', 
        'seasonal_winter_christmas'
    ]
    
    missing_columns = [col for col in seasonal_columns if col not in headers]
    
    if missing_columns:
        print(f"❌ Missing seasonal columns in Excel: {missing_columns}")
        return False
    
    print(f"✓ Excel export contains all 4 seasonal columns")
    
    # Get column indices
    col_indices = {col: headers.index(col) + 1 for col in seasonal_columns}
    
    # Count games with seasonal tags
    seasonal_count = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if any(row[col_indices[col] - 1].value for col in seasonal_columns):
            seasonal_count += 1
    
    print(f"✓ Found {seasonal_count} games with seasonal tags in Excel")
    
    # Show column positions
    print(f"\nSeasonal column positions:")
    for col, idx in col_indices.items():
        print(f"  - {col}: Column {idx}")
    
    return True

def test_text_export():
    """Verify text export does NOT include seasonal data."""
    print("\n=== Testing Text Export ===")
    text_path = Path("reports/games.txt")
    
    if not text_path.exists():
        print(f"❌ Text file not found at {text_path}")
        return False
    
    with open(text_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Text export should only have format: [ID] Title: Score
    lines = [line for line in content.split('\n') if line.strip()]
    
    if not lines:
        print("❌ Text file is empty")
        return False
    
    # Check format of first few lines
    print(f"✓ Text file has {len(lines)} lines")
    print("\nFirst 3 lines:")
    for line in lines[:3]:
        print(f"  {line}")
    
    # Verify no seasonal keywords appear
    seasonal_keywords = ['seasonal_spring', 'seasonal_summer', 'seasonal_fall', 'seasonal_winter']
    found_keywords = [kw for kw in seasonal_keywords if kw in content.lower()]
    
    if found_keywords:
        print(f"❌ Text export contains seasonal keywords: {found_keywords}")
        return False
    
    print("✓ Text export does NOT contain seasonal data (as expected)")
    
    return True

def main():
    """Run all export verification tests."""
    print("=" * 60)
    print("Seasonal Attributes Export Verification")
    print("=" * 60)
    
    results = {
        'JSON': test_json_export(),
        'Excel': test_excel_export(),
        'Text': test_text_export()
    }
    
    print("\n" + "=" * 60)
    print("Summary:")
    print("=" * 60)
    for export_type, passed in results.items():
        status = "✓ PASS" if passed else "❌ FAIL"
        print(f"{export_type}: {status}")
    
    all_passed = all(results.values())
    print("\n" + ("=" * 60))
    if all_passed:
        print("✓ All export tests PASSED")
        print("Requirements 5.1, 5.2, 5.3, 5.4 verified")
    else:
        print("❌ Some export tests FAILED")
    print("=" * 60)
    
    return all_passed

if __name__ == "__main__":
    success = main()
    exit(0 if success else 1)
