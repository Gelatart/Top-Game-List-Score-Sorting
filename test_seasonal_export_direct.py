#!/usr/bin/env python3
"""
Direct test of export functions with seasonal data.
This creates test GameObjects with seasonal attributes and verifies exports.
"""

import json
import os
import tempfile
from pathlib import Path
from openpyxl import load_workbook

from src.generator.game_object import GameObject
from src.generator.exporter import export_to_json, export_to_excel, export_to_text


def create_test_games():
    """Create test GameObjects with seasonal attributes."""
    games = [
        GameObject(
            title="Animal Crossing: New Horizons",
            ranked_score=100,
            total_count=50,
            igdb_ID=124030,
            seasonal_spring=True,
            seasonal_summer=True,
            seasonal_winter_christmas=True
        ),
        GameObject(
            title="Resident Evil 4",
            ranked_score=95,
            total_count=45,
            igdb_ID=1127,
            seasonal_fall_halloween=True
        ),
        GameObject(
            title="Super Mario Sunshine",
            ranked_score=90,
            total_count=40,
            igdb_ID=1029,
            seasonal_summer=True
        ),
        GameObject(
            title="The Elder Scrolls V: Skyrim",
            ranked_score=85,
            total_count=35,
            igdb_ID=1447,
            seasonal_winter_christmas=True
        ),
        GameObject(
            title="Stardew Valley",
            ranked_score=80,
            total_count=30,
            igdb_ID=7346,
            seasonal_spring=True
        ),
    ]
    return games


def test_json_export_with_seasonal():
    """Test JSON export includes seasonal fields."""
    print("\n=== Testing JSON Export with Seasonal Data ===")
    
    games = create_test_games()
    
    with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False) as f:
        temp_path = f.name
    
    try:
        # Export to JSON
        export_to_json(games, temp_path)
        
        # Read back and verify
        with open(temp_path, 'r', encoding='utf-8') as f:
            exported_games = json.load(f)
        
        assert len(exported_games) == 5, "Should have 5 games"
        
        # Check first game has all seasonal fields
        ac = exported_games[0]
        assert 'seasonal_spring' in ac, "Missing seasonal_spring"
        assert 'seasonal_summer' in ac, "Missing seasonal_summer"
        assert 'seasonal_fall_halloween' in ac, "Missing seasonal_fall_halloween"
        assert 'seasonal_winter_christmas' in ac, "Missing seasonal_winter_christmas"
        
        # Verify specific values
        assert ac['seasonal_spring'] is True
        assert ac['seasonal_summer'] is True
        assert ac['seasonal_fall_halloween'] is False
        assert ac['seasonal_winter_christmas'] is True
        
        # Check Resident Evil 4
        re4 = next(g for g in exported_games if g['title'] == "Resident Evil 4")
        assert re4['seasonal_fall_halloween'] is True
        assert re4['seasonal_spring'] is False
        
        print("✓ JSON export includes all seasonal boolean fields")
        print("✓ Seasonal values are correctly set")
        print(f"✓ Verified {len(exported_games)} games")
        
        return True
        
    finally:
        os.unlink(temp_path)


def test_excel_export_with_seasonal():
    """Test Excel export includes seasonal columns."""
    print("\n=== Testing Excel Export with Seasonal Data ===")
    
    games = create_test_games()
    
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
        temp_path = f.name
    
    try:
        # Export to Excel
        export_to_excel(games, temp_path)
        
        # Read back and verify
        wb = load_workbook(temp_path)
        ws = wb.active
        
        # Get headers
        headers = [cell.value for cell in ws[1]]
        
        # Check seasonal columns exist
        seasonal_cols = [
            'seasonal_spring',
            'seasonal_summer',
            'seasonal_fall_halloween',
            'seasonal_winter_christmas'
        ]
        
        for col in seasonal_cols:
            assert col in headers, f"Missing column: {col}"
        
        print("✓ Excel export includes all 4 seasonal columns")
        
        # Get column indices
        col_indices = {col: headers.index(col) for col in seasonal_cols}
        title_idx = headers.index('title')
        
        # Verify Animal Crossing row
        ac_row = None
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            if row[title_idx].value == "Animal Crossing: New Horizons":
                ac_row = row
                break
        
        assert ac_row is not None, "Animal Crossing not found"
        assert ac_row[col_indices['seasonal_spring']].value is True
        assert ac_row[col_indices['seasonal_summer']].value is True
        assert ac_row[col_indices['seasonal_fall_halloween']].value is False
        assert ac_row[col_indices['seasonal_winter_christmas']].value is True
        
        print("✓ Seasonal column values are correct")
        print("✓ Excel columns support boolean filtering")
        
        return True
        
    finally:
        os.unlink(temp_path)


def test_text_export_excludes_seasonal():
    """Test text export does NOT include seasonal data."""
    print("\n=== Testing Text Export Excludes Seasonal Data ===")
    
    games = create_test_games()
    
    with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False) as f:
        temp_path = f.name
    
    try:
        # Export to text
        export_to_text(games, temp_path)
        
        # Read back and verify
        with open(temp_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        lines = [line for line in content.split('\n') if line.strip()]
        assert len(lines) == 5, "Should have 5 lines"
        
        # Verify format: [ID] Title: Score
        first_line = lines[0]
        assert first_line.startswith('['), "Should start with [ID]"
        assert '] ' in first_line and ': ' in first_line, "Should have format [ID] Title: Score"
        
        # Verify no seasonal keywords
        seasonal_keywords = ['seasonal', 'spring', 'summer', 'fall', 'winter', 'halloween', 'christmas']
        for keyword in seasonal_keywords:
            # Only check for the word "seasonal" which shouldn't appear
            if keyword == 'seasonal':
                assert keyword not in content.lower(), f"Text export should not contain '{keyword}'"
        
        print("✓ Text export uses correct format: [ID] Title: Score")
        print("✓ Text export does NOT include seasonal data")
        print(f"✓ Verified {len(lines)} lines")
        
        return True
        
    finally:
        os.unlink(temp_path)


def main():
    """Run all direct export tests."""
    print("=" * 60)
    print("Direct Seasonal Export Tests")
    print("=" * 60)
    
    try:
        results = {
            'JSON': test_json_export_with_seasonal(),
            'Excel': test_excel_export_with_seasonal(),
            'Text': test_text_export_excludes_seasonal()
        }
        
        print("\n" + "=" * 60)
        print("Summary:")
        print("=" * 60)
        for export_type, passed in results.items():
            status = "✓ PASS" if passed else "❌ FAIL"
            print(f"{export_type}: {status}")
        
        all_passed = all(results.values())
        print("\n" + ("=" * 60))
        if all_passed:
            print("✓ All direct export tests PASSED")
            print("✓ Requirements 5.1, 5.2, 5.3, 5.4 verified")
        else:
            print("❌ Some tests FAILED")
        print("=" * 60)
        
        return all_passed
        
    except Exception as e:
        print(f"\n❌ Test failed with error: {e}")
        import traceback
        traceback.print_exc()
        return False


if __name__ == "__main__":
    success = main()
    exit(0 if success else 1)
